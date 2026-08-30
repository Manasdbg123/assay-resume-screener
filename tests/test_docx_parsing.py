"""Round-trip test for the DOCX path, using a document built at test time."""

import io

from docx import Document
from werkzeug.datastructures import FileStorage

from resumescreener.parsing import extract_text


def _docx_bytes():
    document = Document()
    document.add_paragraph("Amara Nwosu - Senior Engineer")
    document.add_paragraph("")  # blank paragraphs must be skipped, not emitted
    document.add_paragraph("8 years of Python and Kubernetes.")

    # Resumes routinely put skills in tables; the parser must read cells too.
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Languages"
    table.cell(0, 1).text = "Python, Go, Rust"

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def test_extracts_paragraphs_and_table_cells():
    text = extract_text(FileStorage(stream=io.BytesIO(_docx_bytes()), filename="cv.docx"))
    assert "Amara Nwosu" in text
    assert "Kubernetes" in text
    assert "Rust" in text          # from the table
    assert "\n\n\n" not in text    # blank paragraphs were skipped
