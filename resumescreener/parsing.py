"""
Resume text extraction for PDF, DOCX, and plain text uploads.

Every failure mode surfaces as `ParsingError` with a message safe to show a user -
the underlying library error is logged, not returned.
"""

import io
import logging

from docx import Document
from pypdf import PdfReader

logger = logging.getLogger(__name__)

# Resumes are a few pages. A 900-page PDF is a resource-exhaustion vector, not a CV.
MAX_PDF_PAGES = 50


class ParsingError(ValueError):
    """A resume could not be read. The message is user-facing."""


def extract_text_from_pdf(file_stream):
    """Extract text from a PDF file stream."""
    try:
        reader = PdfReader(file_stream)
        if len(reader.pages) > MAX_PDF_PAGES:
            raise ParsingError(
                f"This PDF has {len(reader.pages)} pages. Please upload a resume of "
                f"{MAX_PDF_PAGES} pages or fewer."
            )
        text_parts = []
        for page in reader.pages[:MAX_PDF_PAGES]:
            page_text = page.extract_text()
            if page_text:
                text_parts.append(page_text)
        return "\n".join(text_parts)
    except ParsingError:
        raise
    except Exception as exc:
        logger.warning("PDF parsing failed", exc_info=True)
        raise ParsingError(
            "This PDF could not be read. It may be corrupted or password-protected."
        ) from exc


def extract_text_from_docx(file_stream):
    """Extract text from a DOCX file stream."""
    try:
        doc = Document(file_stream)
        text_parts = []
        for paragraph in doc.paragraphs:
            if paragraph.text.strip():
                text_parts.append(paragraph.text)
        # Also extract text from tables
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text.strip():
                        text_parts.append(cell.text)
        return "\n".join(text_parts)
    except Exception as exc:
        logger.warning("DOCX parsing failed", exc_info=True)
        raise ParsingError("This DOCX file could not be read. It may be corrupted.") from exc


def extract_text_from_txt(file_stream):
    """Extract text from a plain text file stream."""
    try:
        if isinstance(file_stream, bytes):
            return file_stream.decode("utf-8", errors="ignore")
        content = file_stream.read()
        if isinstance(content, bytes):
            return content.decode("utf-8", errors="ignore")
        return content
    except Exception as exc:
        logger.warning("text parsing failed", exc_info=True)
        raise ParsingError("This text file could not be read.") from exc


def extract_text(file_storage):
    """
    Route to the appropriate parser based on file extension.

    Args:
        file_storage: Flask FileStorage object from form upload

    Returns:
        Extracted text as a string
    """
    filename = file_storage.filename.lower()
    file_stream = io.BytesIO(file_storage.read())

    if filename.endswith(".pdf"):
        return extract_text_from_pdf(file_stream)
    elif filename.endswith(".docx"):
        return extract_text_from_docx(file_stream)
    elif filename.endswith(".txt"):
        return extract_text_from_txt(file_stream)
    else:
        raise ParsingError(
            "Unsupported file format. Please upload a PDF, DOCX, or TXT file."
        )
